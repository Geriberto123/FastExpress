import os

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "docs/Informe_QA_FastExpress.docx"
BLUE = "17365D"
LIGHT = "EAF2F8"

CFG_PNG = "docs/evidencias/cfg.png"
COBERTURA_PNG = "docs/evidencias/cobertura.png"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        tc_borders.append(element)
    tc_pr.append(tc_borders)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade(cell, BLUE)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
            run.font.size = Pt(9)
    for row_i, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            if row_i % 2 == 1:
                shade(cells[i], LIGHT)
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            borders(cells[i])
    for cell in table.rows[0].cells:
        borders(cell)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph()
    return table


def figura(doc, ruta, titulo, ancho=6.0):
    if not os.path.exists(ruta):
        raise FileNotFoundError(
            f"Falta la figura {ruta}. Genérala antes de ejecutar el script "
            "(ver instrucciones al final del documento de cambios)."
        )
    doc.add_picture(ruta, width=Inches(ancho))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(titulo)
    run.italic = True
    run.font.size = Pt(9)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
for name in ("Heading 1", "Heading 2"):
    styles[name].font.name = "Aptos Display"
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
styles["Heading 1"].font.size = Pt(15)
styles["Heading 2"].font.size = Pt(12)

header = section.header.paragraphs[0]
header.text = "FastExpress  |  Informe de aseguramiento de calidad"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(89, 89, 89)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Proyecto FastExpress  •  Informe QA")
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(89, 89, 89)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Informe de QA FastExpress")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Cotización de envíos y descuentos").italic = True
for linea in (
    "Curso: Calidad de Software",
    "Equipo: Adolfo Andres Cera Urueta, Juan Esteban Ordoñez Tobon, Samuel David Morales Rios",
    "Repositorio: https://github.com/Geriberto123/FastExpress.git",
):
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(linea).font.size = Pt(9)
doc.add_paragraph()

p = doc.add_paragraph()
p.add_run("Resultado ejecutivo. ").bold = True
p.add_run(
    "La funcionalidad de cotización quedó verificada con evidencia automatizada: "
    "177 pruebas de backend, 3 pruebas unitarias de validadores y 11 pruebas de "
    "interfaz superaron la ejecución de la base 7b98b16 registrada en docs/evidencias. Estas cifras "
    "corresponden a esa ejecución; si la versión final añade pruebas, debe repetirse "
    "la ejecución y actualizar este informe."
)

add_table(doc, ["Indicador", "Resultado", "Evidencia"], [
    ["Backend", "177 de 177 pruebas correctas", "docs/evidencias/backend-jest.json"],
    ["Validadores frontend", "3 de 3 correctas", "docs/evidencias/frontend-unit.txt"],
    ["Interfaz Playwright", "11 de 11 correctas", "docs/evidencias/frontend-ui.txt"],
    ["Cobertura", "98.38 % líneas  |  96.96 % ramas", "docs/evidencias/coverage"],
    ["Carga local", "50 de 50 correctas  |  p95 34.15 ms", "docs/evidencias/performance.txt"],
], [1.35, 2.2, 2.9])

heading(doc, "1 Alcance y criterios de aceptación")
doc.add_paragraph(
    "El trabajo cubrió las reglas BR-01 a BR-04: tarifa base de $5.00, tarifa por peso "
    "por tramo, descuento por tipo de cliente y recargo EXPRESS sobre el subtotal antes "
    "del descuento. También se revisaron validaciones, mensajes, API REST, visualización "
    "de la interfaz y respuesta bajo concurrencia."
)
bullet(doc, "Criterio de cobertura: mínimo 85 % en líneas y ramas.")
bullet(doc, "Criterio local de rendimiento: p95 menor de 500 ms y cero respuestas 5xx.")
bullet(doc, "La especificación, no la implementación, se usó como oráculo de resultados.")

heading(doc, "2 Hallazgos y correcciones")
add_table(doc, ["ID", "Hallazgo", "Acción tomada", "Estado"], [
    ["BUG-001", "El HTML referenciaba css/js mientras los directorios versionados son Css/Js. Puede fallar en servidores sensibles a mayúsculas.", "Se corrigieron las rutas de hoja de estilo y scripts.", "Corregido; UI-09 verifica la carga sin 4xx/5xx."],
    ["OBS-001", "Las pruebas de backend tenían anotaciones any que impedían lint limpio.", "Se tiparon o infirieron los parámetros sin cambiar expectativas.", "Cerrada."],
    ["OBS-002", "La pantalla muestra -$0.00 y +$0.00 cuando descuento o recargo son cero.", "UI-03 valida el valor 0.00, sin fijar el signo como requisito.", "Pendiente decisión de producto."],
], [0.7, 2.3, 2.3, 1.15])

heading(doc, "3 Reglas verificadas y resultados")
doc.add_paragraph(
    "Los cálculos se verificaron contra valores independientes del código. A 2.00 kg el "
    "subtotal es $9.00. La matriz siguiente confirma la combinación de descuento y recargo."
)
add_table(doc, ["Cliente", "Envío", "Descuento", "Recargo", "Total", "Caso"], [
    ["NUEVO", "ESTANDAR", "$0.00", "$0.00", "$9.00", "CP-201"],
    ["NUEVO", "EXPRESS", "$0.00", "$2.25", "$11.25", "CP-202"],
    ["FRECUENTE", "ESTANDAR", "$0.90", "$0.00", "$8.10", "CP-203"],
    ["FRECUENTE", "EXPRESS", "$0.90", "$2.25", "$10.35", "CP-204"],
    ["VIP", "ESTANDAR", "$1.80", "$0.00", "$7.20", "CP-205"],
    ["VIP", "EXPRESS", "$1.80", "$2.25", "$9.45", "CP-206"],
    ["INVÁLIDO", "ESTANDAR", "—", "—", "HTTP 400 cliente", "CP-208"],
    ["NUEVO", "INVÁLIDO", "—", "—", "HTTP 400 envío", "CP-209"],
], [1.0, 1.2, 1.0, 1.0, 0.9, 0.8])
doc.add_paragraph(
    "El redondeo se hace en decimal half-up por importe. El caso CP-207 comprueba que "
    "0.03 kg, NUEVO y EXPRESS totaliza $6.33: subtotal $5.06, recargo $1.27 y descuento $0.00."
)

heading(doc, "4 Límites, clases de equivalencia y validación de entradas")
add_table(doc, ["Clase EP", "Ejemplos", "Resultado esperado", "Prueba"], [
    ["Vacío", "''", "El peso es obligatorio", "UI-02"],
    ["No numérico o formato inválido", "abc, 5kg, 1e3, 5,", "Formato inválido", "CP-002"],
    ["Menor o igual a cero", "0, 0.00", "Peso mínimo 0.01 kg", "CP-109"],
    ["Tramo 1", "0.01 a 5.00", "Tarifa de 2 USD/kg", "CP-101 a CP-104"],
    ["Tramo 2", "5.01 a 20.00", "Tarifa de 4 USD/kg", "CP-105 a CP-108"],
    ["Superior al máximo", "20.01", "Límite excedido", "CP-110"],
], [1.65, 1.6, 2.0, 1.4])
doc.add_paragraph("BVA de peso con NUEVO y ESTANDAR:")
add_table(doc, ["Entrada", "Resultado esperado", "Total", "Caso"], [
    ["0.00", "Rechazo por mínimo", "—", "CP-109"],
    ["0.01", "Aceptado", "$5.02", "CP-101"],
    ["5.00", "Aceptado primer tramo", "$15.00", "CP-104"],
    ["5.01", "Aceptado segundo tramo", "$25.04", "CP-105"],
    ["20.00", "Aceptado", "$85.00", "CP-108"],
    ["20.01", "Rechazo por límite", "—", "CP-110"],
    ["20.001", "Formato inválido", "—", "CP-003 / CP-111"],
    ["0.001", "Formato inválido", "—", "CP-002 / CP-112"],
], [1.0, 2.6, 0.9, 1.0])
doc.add_paragraph(
    "Clases de equivalencia del peso: vacío (UI-02), no numérico o formato inválido "
    "(CP-002), menor o igual a cero (CP-109 para 0 y 0.00), tramo 1 (CP-101 a CP-104), "
    "tramo 2 (CP-105 a CP-108) y superior al máximo (CP-110). El valor negativo -1 se "
    "rechaza por formato antes que por mínimo; aún no tiene un caso propio y se recomienda añadirlo."
)
doc.add_paragraph(
    "No existe un campo de entrada llamado monto. En este informe, monto es el total "
    "cotizado de salida; sus límites se prueban mediante los límites del peso y se "
    "resumen en la siguiente tabla."
)
add_table(doc, ["Entrada de la prueba", "Total esperado", "Caso"], [
    ["VIP, 0.01 kg, ESTANDAR (mínimo alcanzable)", "4.02", "CP-301"],
    ["VIP, 0.02 kg, ESTANDAR (vecino superior)", "4.03", "CP-302"],
    ["NUEVO, 0.01 kg, ESTANDAR (referencia)", "5.02", "CP-303"],
    ["NUEVO, 20.00 kg, ESTANDAR (máximo ESTANDAR)", "85.00", "CP-304"],
    ["NUEVO, 20.00 kg, EXPRESS (máximo alcanzable)", "106.25", "CP-305"],
    ["Entradas válidas (no alcanzable)", "4.01 no aparece", "CP-306"],
    ["Entradas válidas (no alcanzable)", "106.26 no aparece", "CP-307"],
], [3.6, 1.6, 1.0])
doc.add_paragraph(
    "CP-306 y CP-307 recorren exhaustivamente los pesos de 0.01 a 20.00 kg y todas las "
    "combinaciones de cliente y envío permitidas; ninguno de los dos totales aparece."
)

heading(doc, "5 Diseño de pruebas y cobertura")
doc.add_paragraph(
    "Se incorporó una suite de caja negra para la API con los CP-101 a CP-307. La interfaz "
    "dispone de UI-01 a UI-09. Las tres pruebas de visualización de la cotización se "
    "distinguen así: UI-08a (acepta coma y muestra el total de la API interceptada), "
    "UI-08b (reproduce una tabla BVA y de decisión con respuestas interceptadas) y UI-08c "
    "(flujo contra el backend real)."
)
add_table(doc, ["Área", "Líneas", "Ramas", "Funciones", "Sentencias"], [
    ["Total", "98.38 %", "96.96 %", "95.83 %", "98.48 %"],
    ["Cotización dominio", "100 %", "100 %", "100 %", "100 %"],
    ["HTTP", "100 %", "96.96 %", "100 %", "100 %"],
    ["Configuración", "100 %", "100 %", "100 %", "100 %"],
], [1.8, 1.1, 1.1, 1.1, 1.2])
figura(doc, COBERTURA_PNG,
       "Figura 2. Captura del reporte de cobertura (lcov-report) de la ejecución medida.")
doc.add_paragraph(
    "No se excluyó ningún archivo de la medición. La rama pendiente de src/http corresponde "
    "a la expresión error.constraints ?? {} dentro de la fábrica de excepciones de validación "
    "(configurar-app.ts, línea 35 en la versión medida y línea 37 en la versión actual); es un "
    "fallback que la validación vigente no alcanza. La rama pendiente de src corresponde a la "
    "condición require.main === module de main.ts, que solo se ejecuta al arrancar el proceso. "
    "Ambas ramas permanecen visibles en el reporte y ninguna se oculta."
)

heading(doc, "6 Caja blanca: algoritmo de cálculo")
doc.add_paragraph(
    "El grafo de flujo de control del cálculo expandido incorpora las seis validaciones de "
    "entrada y la decisión de tramo. Los errores convergen en un único nodo Fin."
)
figura(doc, CFG_PNG,
       "Figura 1. Grafo de flujo de control del cálculo expandido (fuente: docs/evidencias/cfg.mmd).",
       ancho=6.2)
add_table(doc, ["Nodo", "Flujo de control"], [
    ["1", "Inicio"],
    ["2 a 7", "Peso finito, máximo dos decimales, mínimo, máximo, cliente válido y envío válido"],
    ["8", "Decisión de tarifa: peso menor o igual a 5 kg"],
    ["9a / 9b", "Cálculo con tarifa de 2 USD/kg (9a) o de 4 USD/kg (9b)"],
    ["10", "Error de validación"],
    ["11", "Fin único"],
], [1.0, 5.0])
doc.add_paragraph(
    "Complejidad ciclomática del modelo expandido: N = 12 nodos; E = 18 aristas. "
    "V(G) = E − N + 2 = 18 − 12 + 2 = 8, que coincide con siete decisiones más uno. "
    "Los modelos aislados no se suman entre sí: el validador tiene V(G) = 7 y el cálculo "
    "de tarifa, V(G) = 2."
)
add_table(doc, ["Camino", "Nodos", "Entrada y resultado esperado", "Prueba que lo ejerce"], [
    ["P1", "1-2-10-11", "Peso no finito (NaN) → error de peso finito", "Unitaria: validar-entrada.spec.ts"],
    ["P2", "1-2-3-10-11", "Peso 0.001 → error de precisión", "Unitaria: validar-entrada.spec.ts; CP-112 (API)"],
    ["P3", "1-2-3-4-10-11", "Peso 0 → error de mínimo", "Unitaria: validar-entrada.spec.ts; CP-109 (API)"],
    ["P4", "1-2-3-4-5-10-11", "Peso 20.01 → error de máximo", "Unitaria: validar-entrada.spec.ts; CP-110 (API)"],
    ["P5", "1-2-3-4-5-6-10-11", "Cliente inválido (p. ej. 'vip') → error de cliente", "Unitaria: validar-entrada.spec.ts; CP-208 (contrato API)"],
    ["P6", "1-2-3-4-5-6-7-10-11", "Envío inválido (p. ej. 'URGENTE') → error de envío", "Unitaria: validar-entrada.spec.ts; CP-209 (contrato API)"],
    ["P7", "1-…-8-9a-11", "5.00 kg, NUEVO, ESTANDAR → total $15.00", "Unitaria: calculador-tarifa.spec.ts; CP-104"],
    ["P8", "1-…-8-9b-11", "5.01 kg, NUEVO, ESTANDAR → total $25.04", "Unitaria: calculador-tarifa.spec.ts; CP-105"],
], [0.6, 1.6, 2.6, 2.2])
doc.add_paragraph(
    "Las pruebas de API (CP-xxx) verifican el contrato HTTP; la cobertura de las rutas "
    "internas del validador se demuestra con las pruebas unitarias directas. El redondeo "
    "half-up (CP-207) es una prueba de datos sobre P7 y no un camino de control adicional."
)

heading(doc, "7 Rendimiento local")
doc.add_paragraph(
    "Se creó un script Node sin dependencias adicionales que usa fetch y Promise.all para "
    "ejecutar 50 cotizaciones concurrentes. En la ejecución local: 50 solicitudes exitosas, "
    "0 errores 5xx, p50 de 31.74 ms, p95 de 34.15 ms y tiempo total de 65.38 ms. El resultado "
    "cumple el criterio local definido, pero no sustituye una prueba de carga distribuida."
)

heading(doc, "8 Ciclo de defectos y pendientes")
add_table(doc, ["ID", "Estado", "Detalle"], [
    ["BUG-001", "Cerrado", "Carga de recursos del frontend: 404 en /css/styles.css y /js/*.js en servidor sensible a mayúsculas (docs/evidencias/bug-001-repro.txt). Verificado por UI-09."],
    ["OBS-002", "Pendiente de producto", "Presentación de -$0.00 y +$0.00 en cero. No es un defecto confirmado del cálculo."],
], [1.0, 1.6, 4.4])

heading(doc, "9 Conclusiones riesgos y mejoras futuras")
doc.add_paragraph(
    "El proceso de QA de FastExpress concluye con evidencia de las reglas BR-01 a BR-04, "
    "las particiones y valores límite, la tabla de decisión, el redondeo half-up, la API, "
    "la interfaz UI-01 a UI-09, la cobertura y el benchmark local. La ejecución registrada "
    "cumple el umbral de cobertura establecido y no reporta fallos en las suites ejecutadas; "
    "por tanto, el alcance definido para esta entrega final queda verificado."
)
doc.add_paragraph(
    "Para este informe, Monto se documenta y evalúa como el total cotizado de salida. "
    "Esta definición no impide el cierre de la entrega ni modifica los resultados verificados."
)

heading(doc, "10 Evidencias y ejecución")
doc.add_paragraph(
    "Las evidencias se conservan en docs/evidencias. Para reproducir: ejecutar las pruebas "
    "de Backend con cobertura, las pruebas unitarias y Playwright de Frontend, y luego "
    "iniciar Backend en el puerto 8001 para correr tests/performance/cotizacion-local.mjs "
    "con 50 solicitudes concurrentes."
)

doc.save(OUT)
